"""Build a Google-Docs-friendly DOCX from paper2-electronics-submission.md.

Display equations are rendered to PNG (matplotlib mathtext), inline math becomes
Unicode text, tables get borders, and figures are capped to the page width.
Usage: python3 papers/build_docx.py
"""
import re
import subprocess
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC = HERE / "paper2-electronics-submission.md"
OUT = HERE / "paper2-electronics-submission-gdocs.docx"
plt.rcParams["mathtext.fontset"] = "cm"

GREEK = {r"\alpha": "α", r"\beta": "β", r"\delta": "δ", r"\Delta": "Δ", r"\lambda": "λ",
         r"\pi": "π", r"\rho": "ρ", r"\theta": "θ", r"\epsilon": "ε", r"\Sigma": "Σ",
         r"\Phi": "Φ", r"\in": "∈", r"\sim": "∼", r"\cup": "∪", r"\cdot": "·", r"\mid": "|",
         r"\sum": "Σ", r"\lfloor": "⌊", r"\rfloor": "⌋", r"\dots": "…", r"\ast": "*",
         r"\Pr": "Pr", r"\log": "log", r"\max": "max", r"\min": "min", r"\arg": "arg"}
SUB = dict(zip("0123456789aehijklmnoprstuvx+-()", "₀₁₂₃₄₅₆₇₈₉ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ₊₋₍₎"))
SUP = dict(zip("0123456789+-()ni*", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾ⁿⁱ*"))


def braced(s, i):
    """Return (content, next_index) for a {...} group starting at s[i]."""
    depth, j = 0, i
    while j < len(s):
        depth += (s[j] == "{") - (s[j] == "}")
        j += 1
        if depth == 0:
            return s[i + 1:j - 1], j
    return s[i + 1:], len(s)


def to_unicode(m):
    s = m
    s = s.replace("{,}", ",").replace(r"\lbrace", "\0L").replace(r"\rbrace", "\0R")
    s = re.sub(r"\\(?:qquad|quad)", "  ", s)
    s = re.sub(r"\\(?:Big|big|bigg)l?r?", "", s)
    s = re.sub(r"\\[ ,;!]", " ", s)
    out, i = "", 0
    while i < len(s):
        if s[i] == "\\":
            mm = re.match(r"\\[A-Za-z]+", s[i:])
            if not mm:
                out += s[i + 1:i + 2]; i += 2; continue
            cmd = mm.group(0); i += len(cmd)
            if cmd in ("\\mathrm", "\\mathcal", "\\mathbb", "\\text", "\\mathbf"):
                while i < len(s) and s[i] == " ": i += 1
                if i < len(s) and s[i] == "{":
                    inner, i = braced(s, i); out += to_unicode(inner)
            elif cmd in ("\\hat", "\\bar"):
                mark = "\u0302" if cmd == "\\hat" else "\u0304"
                if i < len(s) and s[i] == "{":
                    inner, i = braced(s, i); inner = to_unicode(inner)
                else:
                    t = re.match(r"\\[A-Za-z]+|.", s[i:]); inner = to_unicode(t.group(0)); i += len(t.group(0))
                out += (inner[:1] + mark + inner[1:]) if inner else ""
            elif cmd in ("\\frac", "\\tfrac"):
                a, i = braced(s, i); b, i = braced(s, i)
                out += "(" + to_unicode(a) + ")/(" + to_unicode(b) + ")"
            elif cmd == "\\sqrt":
                a, i = braced(s, i); out += "√(" + to_unicode(a) + ")"
            elif cmd in GREEK:
                out += GREEK[cmd]
            else:
                out += cmd[1:]
        elif s[i] in "_^":
            kind = s[i]; i += 1
            if i < len(s) and s[i] == "{":
                inner, i = braced(s, i)
            else:
                t = re.match(r"\\[A-Za-z]+|.", s[i:])
                inner, i = t.group(0), i + len(t.group(0))
            inner = to_unicode(inner)
            table = SUB if kind == "_" else SUP
            out += "".join(table[c] for c in inner) if inner and all(c in table for c in inner) else kind + inner
        elif s[i] in "{}":
            i += 1
        else:
            out += s[i]; i += 1
    return out.replace("\0L", "{").replace("\0R", "}")


def mathtext(expr):
    e = expr.strip().replace("\n", " ")
    e = re.sub(r"\\(?:Big|big|bigg)[lr]?", "", e)
    e = e.replace(r"\qquad", r"\quad\quad").replace(r"\lbrace", r"\{").replace(r"\rbrace", r"\}")
    e = e.replace(r"\tfrac", r"\frac").replace(r"\dots", r"\ldots").replace(r"\Pr", r"\mathrm{Pr}")
    e = e.replace(r"\arg\max", r"\mathrm{arg\,max}").replace(r"\text", r"\mathrm")
    e = re.sub(r"\\(hat|bar)\s*(\\[A-Za-z]+)", r"\\\1{\2}", e)
    e = re.sub(r"\\(hat|bar)\s*([A-Za-z])", r"\\\1{\2}", e)
    e = e.replace(r"\ /\ ", r"\ /\ ")
    return "$" + e + "$"


def render_equations(md, tmpdir):
    n = [0]

    def repl(m):
        n[0] += 1
        path = Path(tmpdir) / f"eq{n[0]:02d}.png"
        fig = plt.figure(figsize=(0.1, 0.1))
        fig.text(0, 0, mathtext(m.group(1)), fontsize=13)
        fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.06, transparent=False, facecolor="white")
        plt.close(fig)
        return f"\n![]({path})\n"

    return re.sub(r"\$\$(.+?)\$\$", repl, md, flags=re.S)


def set_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)


def main():
    md = SRC.read_text()
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    md = md.replace("../results/plots/paper2/", str(HERE.parent / "results/plots/paper2") + "/")
    with tempfile.TemporaryDirectory() as tmp:
        md = render_equations(md, tmp)
        md = re.sub(r"(?<!\$)\$(?!\$)([^$\n]+?)\$", lambda m: to_unicode(m.group(1)), md)
        tmp_md = Path(tmp) / "paper.md"
        tmp_md.write_text(md)
        subprocess.run(["pandoc", str(tmp_md), "-f", "markdown+pipe_tables-tex_math_dollars", "-o", str(OUT)], check=True)
        doc = Document(str(OUT))
        for t in doc.tables:
            set_borders(t)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(8.5)
        for shape in doc.inline_shapes:
            if shape.width > Inches(6.3):
                ratio = Inches(6.3) / shape.width
                shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)
        doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
